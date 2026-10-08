"""Local evidence extraction (section 6.2).

Parsing runs without macro execution, script evaluation or external link
resolution. Extraction failure is always reported as an actionable failure and
is never silently treated as "no requirements found".
"""
from __future__ import annotations

import csv
import io
import json
import logging
import re
from dataclasses import dataclass, field
from typing import Any

from app.models.common import Warning as DomainWarning
from app.models.source import ExtractionStatus, SourceKind

log = logging.getLogger(__name__)

EXTRACTOR_VERSION = "1.0.0"

DOCUMENT_EXTENSIONS = {
    ".txt": SourceKind.BRIEF, ".md": SourceKind.BRIEF, ".rtf": SourceKind.BRIEF,
    ".docx": SourceKind.WORD, ".doc": SourceKind.WORD,
    ".xlsx": SourceKind.EXCEL, ".xlsm": SourceKind.EXCEL, ".xls": SourceKind.EXCEL,
    ".csv": SourceKind.CSV, ".json": SourceKind.JSON, ".pdf": SourceKind.PDF,
    ".png": SourceKind.IMAGE, ".jpg": SourceKind.IMAGE, ".jpeg": SourceKind.IMAGE,
    ".svg": SourceKind.IMAGE, ".yaml": SourceKind.API_SPEC, ".yml": SourceKind.API_SPEC,
}
RECORDING_EXTENSIONS = {".mp3", ".wav", ".m4a", ".aac", ".ogg", ".flac",
                        ".mp4", ".mov", ".mkv", ".webm", ".avi"}

#: Meeting signals captured alongside plain text (section 6.2).
DECISION_PATTERNS = [
    (r"(?i)^\s*(decision|απόφαση)\s*[:\-]\s*(?P<body>.+)$", "decision"),
    (r"(?i)^\s*(action|todo|ενέργεια)\s*[:\-]\s*(?P<body>.+)$", "action"),
    (r"(?i)^\s*(question|ερώτηση|open question)\s*[:\-]\s*(?P<body>.+)$", "question"),
    (r"(?i)^\s*(requirement|απαίτηση|req)\s*[:\-]\s*(?P<body>.+)$", "requirement"),
    (r"(?i)^\s*(risk|κίνδυνος)\s*[:\-]\s*(?P<body>.+)$", "risk"),
]

TIMESTAMP_RE = re.compile(r"^\[?(\d{1,2}):(\d{2})(?::(\d{2}))?\]?\s*(?:(?P<speaker>[^:]{1,40}):)?\s*(?P<text>.+)$")


@dataclass
class Chunk:
    text: str
    locator_kind: str
    locator_value: str
    page: int | None = None
    start_ms: int | None = None
    end_ms: int | None = None
    speaker: str | None = None
    confidence: float | None = None


@dataclass
class ExtractionResult:
    status: ExtractionStatus
    text: str = ""
    chunks: list[Chunk] = field(default_factory=list)
    meta: dict[str, Any] = field(default_factory=dict)
    warnings: list[DomainWarning] = field(default_factory=list)
    failure_reason: str | None = None
    values_recalculated: bool = False
    signals: list[dict[str, Any]] = field(default_factory=list)
    extractor: str = "text"


def detect_kind(filename: str, declared: SourceKind | None = None) -> SourceKind:
    lowered = filename.lower()
    ext = lowered[lowered.rfind("."):] if "." in lowered else ""
    if ext in RECORDING_EXTENSIONS:
        return SourceKind.RECORDING
    if declared and declared != SourceKind.OTHER:
        return declared
    return DOCUMENT_EXTENSIONS.get(ext, SourceKind.OTHER)


def _decode(data: bytes) -> tuple[str, str]:
    try:
        import chardet

        guess = chardet.detect(data[:200_000]) or {}
        encoding = guess.get("encoding") or "utf-8"
    except Exception:
        encoding = "utf-8"
    for candidate in (encoding, "utf-8", "utf-8-sig", "cp1253", "latin-1"):
        try:
            return data.decode(candidate), candidate
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("utf-8", errors="replace"), "utf-8/replace"


def _collect_signals(text: str, locator: str) -> list[dict[str, Any]]:
    found: list[dict[str, Any]] = []
    for pattern, kind in DECISION_PATTERNS:
        match = re.match(pattern, text)
        if match:
            found.append({"kind": kind, "text": match.group("body").strip(),
                          "locator": locator})
    return found


# --- individual extractors --------------------------------------------------

def _extract_text(data: bytes) -> ExtractionResult:
    text, encoding = _decode(data)
    chunks: list[Chunk] = []
    signals: list[dict[str, Any]] = []
    heading: str | None = None
    for index, block in enumerate(_split_paragraphs(text), start=1):
        if block.startswith("#"):
            heading = block.lstrip("#").strip()
        locator = f"paragraph:{index}" if not heading else f"{heading}#p{index}"
        chunks.append(Chunk(text=block, locator_kind="paragraph", locator_value=locator))
        signals.extend(_collect_signals(block, locator))
    return ExtractionResult(status=ExtractionStatus.READY, text=text, chunks=chunks,
                            meta={"encoding": encoding, "paragraphs": len(chunks)},
                            signals=signals, extractor="text")


def _split_paragraphs(text: str) -> list[str]:
    blocks = [b.strip() for b in re.split(r"\n\s*\n", text)]
    return [b for b in blocks if b]


def _extract_docx(data: bytes) -> ExtractionResult:
    try:
        import docx
    except ImportError:  # pragma: no cover
        return ExtractionResult(status=ExtractionStatus.UNSUPPORTED,
                                failure_reason="python-docx is not installed")
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        return ExtractionResult(
            status=ExtractionStatus.FAILED,
            failure_reason=f"The Word file could not be opened: {exc}. "
                           "Check whether it is password protected or corrupt.",
        )
    chunks: list[Chunk] = []
    lines: list[str] = []
    signals: list[dict[str, Any]] = []
    section = "document"
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower() if paragraph.style else ""
        if style.startswith("heading"):
            section = text
        locator = f"{section}#p{index}"
        chunks.append(Chunk(text=text, locator_kind="paragraph", locator_value=locator))
        lines.append(text)
        signals.extend(_collect_signals(text, locator))
    for t_index, table in enumerate(document.tables, start=1):
        rows: list[str] = []
        for r_index, row in enumerate(table.rows, start=1):
            cells = [c.text.strip() for c in row.cells]
            rows.append(" | ".join(cells))
        body = "\n".join(rows)
        if body.strip():
            locator = f"table:{t_index}"
            chunks.append(Chunk(text=body, locator_kind="table", locator_value=locator))
            lines.append(body)
    warnings: list[DomainWarning] = []
    try:
        has_tracked = b"w:ins" in data or b"w:del" in data
        if has_tracked:
            warnings.append(DomainWarning(
                code="tracked_changes",
                message="The document contains tracked changes. Extracted text reflects "
                        "the stored revision, review the original before relying on it.",
            ))
    except Exception:
        pass
    return ExtractionResult(status=ExtractionStatus.READY, text="\n\n".join(lines),
                            chunks=chunks, signals=signals, warnings=warnings,
                            meta={"paragraphs": len(document.paragraphs),
                                  "tables": len(document.tables)},
                            extractor="python-docx")


def _extract_xlsx(data: bytes) -> ExtractionResult:
    try:
        import openpyxl
    except ImportError:  # pragma: no cover
        return ExtractionResult(status=ExtractionStatus.UNSUPPORTED,
                                failure_reason="openpyxl is not installed")
    try:
        # data_only=True reads the values cached by the producing application.
        # Formulas are never evaluated here.
        book = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        formula_book = openpyxl.load_workbook(io.BytesIO(data), data_only=False, read_only=True)
    except Exception as exc:
        return ExtractionResult(
            status=ExtractionStatus.FAILED,
            failure_reason=f"The workbook could not be opened: {exc}. "
                           "Password protected or corrupt files must be fixed before import.",
        )
    chunks: list[Chunk] = []
    lines: list[str] = []
    sheets: list[dict[str, Any]] = []
    formula_cells = 0
    for sheet in book.worksheets:
        formula_sheet = formula_book[sheet.title] if sheet.title in formula_book.sheetnames else None
        rows: list[str] = []
        max_row = 0
        for r_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            values = ["" if v is None else str(v) for v in row]
            if not any(v.strip() for v in values):
                continue
            max_row = r_index
            rows.append(" | ".join(values))
            if r_index <= 2000:
                locator = f"{sheet.title}!A{r_index}:{_col_letter(len(values))}{r_index}"
                chunks.append(Chunk(text=" | ".join(values), locator_kind="cell_range",
                                    locator_value=locator))
        if formula_sheet is not None:
            try:
                for frow in formula_sheet.iter_rows(values_only=True):
                    formula_cells += sum(
                        1 for v in frow if isinstance(v, str) and v.startswith("=")
                    )
            except Exception:
                pass
        sheets.append({"name": sheet.title, "rows": max_row})
        lines.append(f"## Sheet {sheet.title}\n" + "\n".join(rows))
    warnings = [DomainWarning(
        code="cached_formula_values",
        message=f"{formula_cells} formula cells were read as values cached by the "
                "producing application. Values were not recalculated here and may be stale.",
    )] if formula_cells else []
    return ExtractionResult(status=ExtractionStatus.READY, text="\n\n".join(lines),
                            chunks=chunks, warnings=warnings,
                            meta={"sheets": sheets, "formulaCells": formula_cells},
                            values_recalculated=False, extractor="openpyxl")


def _extract_xls(data: bytes) -> ExtractionResult:
    try:
        import xlrd
    except ImportError:  # pragma: no cover
        return ExtractionResult(status=ExtractionStatus.UNSUPPORTED,
                                failure_reason="xlrd is not installed")
    try:
        book = xlrd.open_workbook(file_contents=data)
    except Exception as exc:
        return ExtractionResult(status=ExtractionStatus.FAILED,
                                failure_reason=f"The legacy workbook could not be opened: {exc}")
    chunks: list[Chunk] = []
    lines: list[str] = []
    sheets = []
    for sheet in book.sheets():
        rows = []
        for r_index in range(sheet.nrows):
            values = [str(sheet.cell_value(r_index, c)) for c in range(sheet.ncols)]
            if not any(v.strip() for v in values):
                continue
            joined = " | ".join(values)
            rows.append(joined)
            locator = f"{sheet.name}!A{r_index + 1}:{_col_letter(sheet.ncols)}{r_index + 1}"
            chunks.append(Chunk(text=joined, locator_kind="cell_range", locator_value=locator))
        sheets.append({"name": sheet.name, "rows": sheet.nrows})
        lines.append(f"## Sheet {sheet.name}\n" + "\n".join(rows))
    return ExtractionResult(
        status=ExtractionStatus.READY, text="\n\n".join(lines), chunks=chunks,
        meta={"sheets": sheets},
        warnings=[DomainWarning(code="cached_formula_values",
                                message="Legacy .xls values are read as stored. "
                                        "They were not recalculated.")],
        extractor="xlrd")


def _col_letter(index: int) -> str:
    index = max(index, 1)
    letters = ""
    while index:
        index, rem = divmod(index - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def _extract_csv(data: bytes) -> ExtractionResult:
    text, encoding = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096])
        delimiter = dialect.delimiter
    except Exception:
        delimiter = ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = list(reader)
    header = rows[0] if rows else []
    chunks = []
    for index, row in enumerate(rows[1:], start=2):
        if not any(cell.strip() for cell in row):
            continue
        pairs = " | ".join(f"{h}={v}" for h, v in zip(header, row)) if header else " | ".join(row)
        chunks.append(Chunk(text=pairs, locator_kind="row", locator_value=f"row:{index}"))
    return ExtractionResult(status=ExtractionStatus.READY, text=text, chunks=chunks,
                            meta={"encoding": encoding, "delimiter": delimiter,
                                  "columns": header, "rows": max(len(rows) - 1, 0)},
                            extractor="csv")


def _extract_json(data: bytes) -> ExtractionResult:
    text, encoding = _decode(data)
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError as exc:
        return ExtractionResult(status=ExtractionStatus.FAILED,
                                failure_reason=f"Invalid JSON at line {exc.lineno}: {exc.msg}")
    chunks: list[Chunk] = []

    def walk(node: Any, path: str) -> None:
        if isinstance(node, dict):
            for key, value in node.items():
                walk(value, f"{path}.{key}" if path else key)
        elif isinstance(node, list):
            for index, value in enumerate(node):
                walk(value, f"{path}[{index}]")
        else:
            chunks.append(Chunk(text=f"{path} = {node}", locator_kind="schema_path",
                                locator_value=path or "$"))

    walk(parsed, "")
    is_openapi = isinstance(parsed, dict) and ("openapi" in parsed or "swagger" in parsed)
    return ExtractionResult(status=ExtractionStatus.READY, text=text, chunks=chunks[:5000],
                            meta={"encoding": encoding, "openapi": is_openapi,
                                  "leafNodes": len(chunks)},
                            extractor="json")


def _extract_pdf(data: bytes) -> ExtractionResult:
    try:
        from pypdf import PdfReader
    except ImportError:  # pragma: no cover
        return ExtractionResult(status=ExtractionStatus.UNSUPPORTED,
                                failure_reason="pypdf is not installed")
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                return ExtractionResult(
                    status=ExtractionStatus.NEEDS_PASSWORD,
                    failure_reason="The PDF is password protected. Supply an unprotected "
                                   "copy or the password to continue.")
    except Exception as exc:
        return ExtractionResult(status=ExtractionStatus.FAILED,
                                failure_reason=f"The PDF could not be opened: {exc}")
    chunks: list[Chunk] = []
    lines: list[str] = []
    empty_pages = 0
    for page_number, page in enumerate(reader.pages, start=1):
        try:
            page_text = page.extract_text() or ""
        except Exception:
            page_text = ""
        if not page_text.strip():
            empty_pages += 1
            continue
        lines.append(page_text)
        for index, block in enumerate(_split_paragraphs(page_text), start=1):
            chunks.append(Chunk(text=block, locator_kind="page",
                                locator_value=f"page:{page_number}#b{index}",
                                page=page_number))
    warnings: list[DomainWarning] = []
    if empty_pages:
        warnings.append(DomainWarning(
            code="no_text_layer",
            message=f"{empty_pages} page(s) contain no extractable text layer. "
                    "They need local OCR before their content can be analysed.",
        ))
    status = ExtractionStatus.READY if chunks else ExtractionStatus.FAILED
    failure = None if chunks else ("No text could be extracted from the PDF. "
                                   "It is probably a scan requiring local OCR.")
    return ExtractionResult(status=status, text="\n\n".join(lines), chunks=chunks,
                            warnings=warnings, failure_reason=failure,
                            meta={"pages": len(reader.pages), "emptyPages": empty_pages},
                            extractor="pypdf")


def _extract_image(_: bytes) -> ExtractionResult:
    """Images need a local OCR engine. The absence of one is reported honestly."""
    return ExtractionResult(
        status=ExtractionStatus.UNSUPPORTED,
        failure_reason="No local OCR engine is configured in this deployment. "
                       "Install the OCR worker, or attach a text export of the image. "
                       "The file is stored and its content is not treated as analysed.",
        warnings=[DomainWarning(code="ocr_required",
                                message="Design screenshots yield observed UI facts only; "
                                        "interaction and business behaviour must be asked about.")],
        extractor="none",
    )


def _extract_recording(data: bytes, filename: str) -> ExtractionResult:
    """Local transcription runs in the AI worker.

    The upload path only records media metadata and queues the job; it never
    claims a transcript exists. A sidecar transcript (``name.transcript.txt``
    style content supplied as the upload) is parsed when timestamps are present.
    """
    return ExtractionResult(
        status=ExtractionStatus.PENDING,
        meta={"filename": filename, "bytes": len(data), "requiresTranscription": True},
        warnings=[DomainWarning(
            code="transcription_queued",
            message="Local transcription and diarization are queued. Speaker labels are "
                    "suggestions and never imply verified speaker identity.")],
        extractor="queued",
    )


def parse_transcript_text(text: str) -> list[Chunk]:
    """Parse ``[mm:ss] Speaker: text`` transcripts into timestamped chunks."""
    chunks: list[Chunk] = []
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        match = TIMESTAMP_RE.match(line)
        if not match:
            continue
        hours_or_minutes = int(match.group(1))
        minutes_or_seconds = int(match.group(2))
        seconds = match.group(3)
        if seconds is not None:
            start_ms = (hours_or_minutes * 3600 + minutes_or_seconds * 60 + int(seconds)) * 1000
        else:
            start_ms = (hours_or_minutes * 60 + minutes_or_seconds) * 1000
        chunks.append(Chunk(
            text=match.group("text").strip(),
            locator_kind="timestamp",
            locator_value=f"t:{start_ms}",
            start_ms=start_ms,
            end_ms=start_ms + 5000,
            speaker=(match.group("speaker") or "").strip() or None,
        ))
    return chunks


def extract(data: bytes, filename: str, kind: SourceKind) -> ExtractionResult:
    """Dispatch to the extractor for a file kind. Never raises."""
    lowered = filename.lower()
    try:
        if kind == SourceKind.RECORDING:
            return _extract_recording(data, filename)
        if lowered.endswith(".docx"):
            return _extract_docx(data)
        if lowered.endswith(".doc"):
            return ExtractionResult(
                status=ExtractionStatus.UNSUPPORTED,
                failure_reason="Legacy .doc needs the isolated conversion worker. "
                               "Convert to .docx or enable the converter service.")
        if lowered.endswith((".xlsx", ".xlsm")):
            return _extract_xlsx(data)
        if lowered.endswith(".xls"):
            return _extract_xls(data)
        if lowered.endswith(".csv"):
            return _extract_csv(data)
        if lowered.endswith(".json"):
            return _extract_json(data)
        if lowered.endswith(".pdf"):
            return _extract_pdf(data)
        if lowered.endswith((".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp")):
            return _extract_image(data)
        if lowered.endswith((".txt", ".md", ".markdown", ".yaml", ".yml", ".feature",
                             ".log", ".rtf", ".xml", ".html")):
            result = _extract_text(data)
            transcript_chunks = parse_transcript_text(result.text)
            if len(transcript_chunks) >= 3:
                result.chunks = transcript_chunks
                result.meta["transcript"] = True
            return result
        # Unknown binary: store it, do not guess at its content.
        return ExtractionResult(
            status=ExtractionStatus.UNSUPPORTED,
            failure_reason=f"No extractor is registered for '{filename}'. The original is "
                           "stored and must be summarised manually or converted.")
    except Exception as exc:  # pragma: no cover - defensive
        log.exception("extraction failed for %s", filename)
        return ExtractionResult(status=ExtractionStatus.FAILED,
                                failure_reason=f"Extraction error: {exc}")
